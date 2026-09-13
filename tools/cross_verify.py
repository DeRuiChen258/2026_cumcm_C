"""独立复核：用 openpyxl 只读地把原始附件、clean 数据和五个结果文件对一遍。

跑法（要装了 openpyxl 的 Python）：
    /usr/bin/python3 tools/cross_verify.py
它不在主流水线里，属于另开一条通道的交叉检查。
"""
from __future__ import annotations

import ast
import json
import random
import struct
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT: dict = {"checks": [], "all_pass": True}


def record(name: str, ok: bool, detail: str) -> None:
    OUT["checks"].append({"name": name, "pass": bool(ok), "detail": detail})
    OUT["all_pass"] = OUT["all_pass"] and bool(ok)


def read_npy(path: Path):
    """Minimal C-order .npy reader (float64 / int64), stdlib only."""
    raw = path.read_bytes()
    if raw[:6] != b"\x93NUMPY":
        raise ValueError(f"{path} is not a .npy file")
    header_len = struct.unpack("<H", raw[8:10])[0]
    header = ast.literal_eval(raw[10:10 + header_len].decode("latin1").strip())
    start = 10 + header_len
    fmt = {"<f8": "d", "<i8": "q"}[header["descr"]]
    count = 1
    for dim in header["shape"]:
        count *= dim
    values = struct.unpack("<" + fmt * count, raw[start:start + count * struct.calcsize(fmt)])
    return header["shape"], list(values)


def main() -> None:
    # --- 1. slot mapping: clean matrices vs raw attachment cells -----------------
    shape, load = read_npy(ROOT / "clean" / "load.npy")
    _, pv = read_npy(ROOT / "clean" / "pv_actual_raw.npy")
    _, price = read_npy(ROOT / "clean" / "price.npy")
    att2 = openpyxl.load_workbook(ROOT / "Data" / "附件2.xlsx", data_only=True)
    att4 = openpyxl.load_workbook(ROOT / "Data" / "附件4.xlsx", data_only=True)
    load_ws = att2["小区负载"]
    pv_ws = att2["光伏发电实际功率"]
    price_ws = att4[att4.sheetnames[0]]
    rng = random.Random(20250911)
    samples = [(rng.randrange(shape[0]), rng.randrange(shape[1])) for _ in range(3)]
    for day, k in samples:
        row, col = day + 2, k + 2
        ok = (abs(load[day * 144 + k] - float(load_ws.cell(row, col).value)) < 1e-9
              and abs(pv[day * 144 + k] - float(pv_ws.cell(row, col).value)) < 1e-9
              and abs(price[day * 144 + k] - float(price_ws.cell(row, col).value)) < 1e-9)
        record(f"slot mapping day {day} slot {k+1}", ok,
               f"clean=({load[day*144+k]:.4f},{pv[day*144+k]:.4f},{price[day*144+k]:.4f}) vs "
               f"attachment row {row} col {col}")

    # --- 2. wide-table values vs the stored arrays ------------------------------
    _, q2_plan = read_npy(ROOT / "results" / "arrays" / "q2_q_plan.npy")
    wb = openpyxl.load_workbook(ROOT / "output" / "result2.xlsx", data_only=True)
    ws = wb["计划购电量"]
    mismatches = 0
    for i in range(0, 334, 37):           # spot-check 10 rows spread over the year
        for k in range(0, 144, 17):       # and 9 slots per row
            if abs(float(ws.cell(i + 2, k + 2).value) - q2_plan[i * 144 + k]) > 1e-6:
                mismatches += 1
    record("result2 wide table matches stored plan array", mismatches == 0,
           f"{mismatches} mismatching spot-checked cells (90 samples)")

    # --- 3. template fidelity ---------------------------------------------------
    for name in ("result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"):
        tpl = openpyxl.load_workbook(ROOT / "templates" / "附件5" / name)
        cur = openpyxl.load_workbook(ROOT / "output" / name)
        same_sheets = tpl.sheetnames == cur.sheetnames
        header_ok = True
        for sheet in tpl.sheetnames:
            a, b = tpl[sheet], cur[sheet]
            for col in range(1, a.max_column + 1):
                if a.cell(1, col).value != b.cell(1, col).value:
                    header_ok = False
        record(f"{name} sheet names and headers identical to attachment-5 template",
               same_sheets and header_ok, f"sheets={cur.sheetnames}")

    # --- 4. dates remain dates and the daily numbers add up ---------------------
    ws = wb["计划购电量"]
    date_ok = all(hasattr(ws.cell(r, 1).value, "year") for r in (2, 100, 335))
    record("result2 dates are Excel date cells", date_ok,
           str([ws.cell(r, 1).value for r in (2, 100, 335)]))
    total_ok = True
    for r in (2, 150, 335):
        row_sum = sum(float(ws.cell(r, c).value) for c in range(2, 146))
        total_ok = total_ok and abs(row_sum - float(ws.cell(r, 146).value)) < 1e-6
    record("daily purchase column equals the sum of the 144 slots", total_ok,
           "checked rows 2 / 150 / 335")
    (ROOT / "results" / "cross_verify_xlsx.json").write_text(
        json.dumps(OUT, ensure_ascii=False, indent=2), encoding="utf-8")
    for check in OUT["checks"]:
        print(f"[{'PASS' if check['pass'] else 'FAIL'}] {check['name']} :: {check['detail']}")
    print(f"[cross-verify] all_pass = {OUT['all_pass']}")


if __name__ == "__main__":
    main()
