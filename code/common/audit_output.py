"""独立复算 `output/result*.xlsx`：以结果文件为真相源的账目与物理一致性审计。

设计原则（只读）：

* 读取通道与写入通道分离：用 openpyxl 读 xlsx（写入端用自实现的 zip+XML），
  避免"写入端自证"。
* 一切判据都回到题面口径：Q1/Q2/Q3 按附件 1 电价结算，Q4-2/Q4-3 按附件 4 电价结算；
  紧急购电按 5 倍交易时刻电价；Q3/Q4-3 的计划—调整差按 0.5 / 1.5 倍结算。
* 输出 `results/output_audit.json`：逐项 PASS/FAIL、最大偏差与论文引用的关键数字。

用法：
    python3 -m code.common.audit_output
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import numpy as np
import openpyxl

from code.common.load_clean import ROOT, load_panel

RESULTS = ROOT / "results"
OUTPUT = ROOT / "output"
TEMPLATES = ROOT / "templates" / "附件5"
FIRST_DAY = 31
BLOCK_RE = re.compile(r"^(\d+):(\d+)-(\d+):(\d+)$")


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def block_slots(label: str) -> tuple[int, int]:
    """把 '7:20-8:10' 映射为覆盖的 10 分钟时段区间 [k0, k1)（右端点口径）。"""
    m = BLOCK_RE.match(str(label).strip())
    if m is None:
        raise ValueError(f"unparsable block label: {label!r}")
    h0, m0, h1, m1 = (int(g) for g in m.groups())
    start, end = h0 * 60 + m0, h1 * 60 + m1
    if end <= start:
        raise ValueError(f"non-positive block span: {label!r}")
    return start // 10, end // 10


def sheet_rows(ws) -> list[tuple]:
    return [row for row in ws.iter_rows(values_only=True)]


def read_slots_sheet(ws) -> tuple[list[str], np.ndarray]:
    rows = sheet_rows(ws)
    labels = [str(r[0]) for r in rows[1:]]
    values = np.array([float(r[1]) if r[1] is not None else np.nan for r in rows[1:]], dtype=float)
    return labels, values


def read_wide_sheet(ws) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    rows = sheet_rows(ws)
    values = np.array([[float(v) for v in r[1:145]] for r in rows[1:]], dtype=float)
    daily_energy = np.array([float(r[145]) for r in rows[1:]], dtype=float)
    daily_cost = np.array([float(r[146]) for r in rows[1:]], dtype=float)
    return values, daily_energy, daily_cost


def read_charge_rows(ws) -> dict[str, np.ndarray]:
    rows = sheet_rows(ws)
    has_date = str(rows[0][0]).strip() == "日期"
    off = 1 if has_date else 0
    charge, discharge, soc = [], [], []
    current_charge = current_discharge = 0.0
    soc_start = soc_end = None
    for row in rows[1:]:
        if has_date and row[0] is not None:
            if soc_start is not None:
                charge.append(current_charge)
                discharge.append(current_discharge)
                soc.append((soc_start, soc_end))
            current_charge = current_discharge = 0.0
            soc_start = float(row[off + 4])
            soc_end = None
        if not has_date and soc_start is None:
            soc_start = float(row[off + 4])
        current_charge += float(row[off + 1])
        current_discharge += float(row[off + 2])
        if row[off + 3] is not None and str(row[off + 3]).strip() == "24:00":
            soc_end = float(row[off + 4])
    charge.append(current_charge)
    discharge.append(current_discharge)
    soc.append((soc_start, soc_end))
    return {"charge": np.array(charge), "discharge": np.array(discharge),
            "soc_start": np.array([s[0] for s in soc], dtype=float),
            "soc_end": np.array([s[1] for s in soc], dtype=float)}


def read_emergency_rows(ws) -> tuple[np.ndarray, list[dict]]:
    """返回每日紧急购电量（按行出现顺序累加）与原始块记录。"""
    rows = sheet_rows(ws)
    energy: list[float] = []
    blocks: list[dict] = []
    for row in rows[1:]:
        if row[0] is not None:
            energy.append(0.0)
        if row[1] is None and row[2] is None:
            continue
        value = float(row[2])
        energy[-1] += value
        blocks.append({"day": len(energy) - 1, "label": str(row[1]), "energy_kwh": value})
    return np.array(energy, dtype=float), blocks


def price_for(tag: str, panel) -> np.ndarray:
    """Q2/Q3 按附件 1 电价（题目："每天的电价相同"），Q4 按附件 4 电价。"""
    if tag.startswith("q4"):
        return panel.price
    return np.tile(panel.day1[:, 0], (panel.price.shape[0], 1))


def add(checks: list[dict], name: str, ok: bool, detail: str = "", dev: float | None = None) -> None:
    checks.append({"check": name, "pass": bool(ok), "detail": detail,
                   "max_abs_deviation": None if dev is None else float(dev)})


def audit_result1(panel, checks: list[dict]) -> dict:
    st = panel.config["storage"]
    eta_c, eta_d = float(st["eta_charge"]), float(st["eta_discharge"])
    soc_init = float(st["soc_init_kwh"])
    wb = openpyxl.load_workbook(OUTPUT / "result1.xlsx", read_only=True, data_only=True)
    tmpl = openpyxl.load_workbook(TEMPLATES / "result1.xlsx", read_only=True, data_only=True)
    labels, q = read_slots_sheet(wb["计划购电量"])
    tmpl_labels, _ = read_slots_sheet(tmpl["计划购电量"])
    add(checks, "result1 时段标签与模板逐字一致", labels == tmpl_labels,
        f"{len(labels)} 行")
    cd = read_charge_rows(wb["充放电量"])
    wb.close()
    tmpl.close()

    q1 = load_json(RESULTS / "q1.json")
    price = panel.day1[:, 0]
    cost = float(np.sum(price * q))
    add(checks, "result1 计划购电量合计 = q1.json 计划量",
        abs(q.sum() - float(np.sum(q1["lp"]["q"]))) < 1e-6,
        f"{q.sum():.2f} vs {float(np.sum(q1['lp']['q'])):.2f}")
    add(checks, "result1 购电费 = Σ 附件1电价×购电量 = q1.json 最优值",
        abs(cost - float(q1["cost_cny"])) < 1e-6,
        f"{cost:.6f} vs {float(q1['cost_cny']):.6f}", abs(cost - float(q1["cost_cny"])))
    add(checks, "result1 0:00 与 24:00 储电量均等于初始 6000 kWh",
        abs(cd["soc_start"][0] - soc_init) < 1e-9 and abs(cd["soc_end"][0] - soc_init) < 1e-9,
        f"{cd['soc_start'][0]:.1f} / {cd['soc_end'][0]:.1f}")
    delta = eta_c * cd["charge"][0] - cd["discharge"][0] / eta_d
    add(checks, "result1 充放电量与 SOC 日循环恒等式",
        abs(delta) < 1e-6, f"Σ(ηc·c − d/ηd) = {delta:.9f} kWh", abs(delta))
    add(checks, "result1 充放电量与 q1.json 一致",
        abs(cd["charge"][0] - float(np.sum(q1["lp"]["c"]))) < 1e-6
        and abs(cd["discharge"][0] - float(np.sum(q1["lp"]["d"]))) < 1e-6,
        f"{cd['charge'][0]:.2f} / {cd['discharge'][0]:.2f}")
    return {
        "purchase_kwh": float(q.sum()),
        "cost_cny": cost,
        "charge_kwh": float(cd["charge"][0]),
        "discharge_kwh": float(cd["discharge"][0]),
        "soc_start_kwh": float(cd["soc_start"][0]),
        "soc_end_kwh": float(cd["soc_end"][0]),
    }


def audit_wide(tag: str, file: str, with_adjust: bool, panel, checks: list[dict]) -> dict:
    st = panel.config["storage"]
    eta_c, eta_d = float(st["eta_charge"]), float(st["eta_discharge"])
    soc_init = float(st["soc_init_kwh"])
    days = slice(FIRST_DAY, 365)
    price = price_for(tag, panel)[days]
    summary = load_json(RESULTS / f"{tag}.json")["summary"]
    records = {r["date"]: r for r in load_json(RESULTS / f"{tag}.json")["records"]}

    wb = openpyxl.load_workbook(OUTPUT / file, read_only=True, data_only=True)
    plan, plan_daily, plan_cost = read_wide_sheet(wb["计划购电量"])
    adj = plan.copy()
    if with_adjust:
        adj, adj_daily, _ = read_wide_sheet(wb["调整购电量"])
        add(checks, f"{tag} 调整购电量日合计列自洽",
            float(np.abs(adj_daily - adj.sum(axis=1)).max()) < 1e-6,
            "", float(np.abs(adj_daily - adj.sum(axis=1)).max()))
    emergency, blocks = read_emergency_rows(wb["紧急购电量"])
    cd = read_charge_rows(wb["充放电量"])
    wb.close()

    add(checks, f"{tag} 计划购电量日合计列自洽",
        float(np.abs(plan_daily - plan.sum(axis=1)).max()) < 1e-6,
        "", float(np.abs(plan_daily - plan.sum(axis=1)).max()))
    dev_cost = float(np.abs(np.sum(price * plan, axis=1) - plan_cost).max())
    add(checks, f"{tag} 逐日购电费列 = Σ 电价×计划量（{'附件4' if tag.startswith('q4') else '附件1'}电价）",
        dev_cost < 1e-6, "", dev_cost)

    plan_cost_total = float(np.sum(plan_cost))
    emergency_energy = float(emergency.sum())
    add(checks, f"{tag} 紧急购电量合计 = 记录值",
        abs(emergency_energy - float(summary["total_emergency_energy_kwh"])) < 1e-6,
        f"{emergency_energy:.6f} vs {float(summary['total_emergency_energy_kwh']):.6f}",
        abs(emergency_energy - float(summary["total_emergency_energy_kwh"])))

    # 紧急购电块按覆盖时段均匀分摊，用于把 5 倍费用摊回时段口径后与记录值比对
    g_slot = np.zeros_like(plan)
    for block in blocks:
        k0, k1 = block_slots(block["label"])
        g_slot[block["day"], k0:k1] += block["energy_kwh"] / (k1 - k0)
    emergency_cost = float(5.0 * np.sum(price * g_slot))

    over = float(np.sum(price * np.maximum(0.0, plan - adj)))
    under = float(np.sum(price * np.maximum(0.0, adj - plan)))
    over_pen = 0.5 * over
    under_pen = 1.5 * under
    if with_adjust:
        add(checks, f"{tag} 超额违约费 0.5Σπ(q^plan−q^adj)^+ = 记录值",
            abs(over_pen - float(summary["total_over_plan_penalty_cny"])) < 1e-6,
            f"{over_pen:.6f} vs {float(summary['total_over_plan_penalty_cny']):.6f}",
            abs(over_pen - float(summary["total_over_plan_penalty_cny"])))
        add(checks, f"{tag} 欠额违约费 1.5Σπ(q^adj−q^plan)^+ = 记录值",
            abs(under_pen - float(summary["total_under_plan_penalty_cny"])) < 1e-6,
            f"{under_pen:.6f} vs {float(summary['total_under_plan_penalty_cny']):.6f}",
            abs(under_pen - float(summary["total_under_plan_penalty_cny"])))
    else:
        over_pen = under_pen = 0.0

    total = plan_cost_total + float(summary["total_emergency_cost_cny"]) + over_pen + under_pen
    add(checks, f"{tag} 账目合计 = 计划费+紧急费+违约费 = 记录总额",
        abs(total - float(summary["total_cost_cny"])) < 1e-6,
        f"{total:.6f} vs {float(summary['total_cost_cny']):.6f}",
        abs(total - float(summary["total_cost_cny"])))
    add(checks, f"{tag} 紧急购电费（5 倍）与记录值同量级（块内均匀分摊口径）",
        abs(emergency_cost - float(summary["total_emergency_cost_cny"]))
        < 0.05 * float(summary["total_emergency_cost_cny"]),
        f"{emergency_cost:.2f} vs {float(summary['total_emergency_cost_cny']):.2f}")

    add(checks, f"{tag} 每日 0:00 与 24:00 储电量均为 6000 kWh",
        float(np.abs(cd["soc_start"] - soc_init).max()) < 1e-9
        and float(np.abs(cd["soc_end"] - soc_init).max()) < 1e-9,
        f"max|Δ| = {max(float(np.abs(cd['soc_start'] - soc_init).max()), float(np.abs(cd['soc_end'] - soc_init).max())):.3e}")
    delta = eta_c * cd["charge"] - cd["discharge"] / eta_d
    add(checks, f"{tag} 充放电量与 SOC 日循环恒等式",
        float(np.abs(delta).max()) < 1e-6, "", float(np.abs(delta).max()))
    add(checks, f"{tag} 单日充放电量不超过 6 h 额定功率",
        float(cd["charge"].max()) <= 6 * float(st["power_kw"]) + 1e-6
        and float(cd["discharge"].max()) <= 6 * float(st["power_kw"]) + 1e-6,
        f"max c={cd['charge'].max():.1f} / max d={cd['discharge'].max():.1f}")

    # 日级功率平衡：光伏 + 执行购电 + 放电 + 紧急 ≥ 负荷 + 充电
    # 注：充放电工作表按 4 h 分块汇总，故只能做日级（总量）核对。
    pv = panel.pv_filled()[days]
    load = panel.load[days]
    dt = float(panel.config["time"]["dt_hours"])
    balance = (np.sum(pv, axis=1) * dt + np.sum(adj, axis=1) + cd["discharge"]
               - np.sum(load, axis=1) * dt - cd["charge"] + emergency)
    add(checks, f"{tag} 日级功率平衡不违反（光伏+购电+放电+紧急 ≥ 负荷+充电）",
        float(balance.min()) > -1e-6, f"最小余量 {balance.min():.6f} kWh", float(balance.min()))

    dates = list(panel.dates[days])
    plan_key = "q_plan_kwh" if "q_plan_kwh" in records[dates[0]] else "planned_purchase_kwh"
    rec_plan = np.array([records[d][plan_key] for d in dates])
    add(checks, f"{tag} 逐日计划购电量 = 记录值",
        float(np.abs(plan.sum(axis=1) - rec_plan).max()) < 1e-4,
        "", float(np.abs(plan.sum(axis=1) - rec_plan).max()))
    rec_charge = np.array([records[d]["charge_kwh"] for d in dates])
    rec_discharge = np.array([records[d]["discharge_kwh"] for d in dates])
    add(checks, f"{tag} 逐日充放电量 = 记录值",
        float(np.abs(cd["charge"] - rec_charge).max()) < 1e-4
        and float(np.abs(cd["discharge"] - rec_discharge).max()) < 1e-4,
        "", max(float(np.abs(cd["charge"] - rec_charge).max()),
                float(np.abs(cd["discharge"] - rec_discharge).max())))
    rec_g = np.array([records[d]["emergency_energy_kwh"] for d in dates])
    add(checks, f"{tag} 逐日紧急购电量 = 记录值",
        float(np.abs(emergency - rec_g).max()) < 1e-4, "",
        float(np.abs(emergency - rec_g).max()))

    return {
        "days": int(plan.shape[0]),
        "total_cost_cny": float(summary["total_cost_cny"]),
        "plan_cost_cny": plan_cost_total,
        "emergency_cost_cny": float(summary["total_emergency_cost_cny"]),
        "emergency_energy_kwh": emergency_energy,
        "over_plan_penalty_cny": over_pen,
        "under_plan_penalty_cny": under_pen,
        "purchase_kwh": float(plan.sum()),
        "adjust_kwh": float(adj.sum()),
        "emergency_blocks": len(blocks),
        "min_daily_balance_kwh": float(balance.min()),
    }


def main() -> None:
    panel = load_panel()
    checks: list[dict] = []
    out = {
        "source": "output/result{1,2,3,4-2,4-3}.xlsx",
        "convention": {
            "q1": "附件1 电价（元/kWh）× 计划购电量",
            "q2_q3": "附件1 电价（题面：每天的电价相同）",
            "q4_2_q4_3": "附件4 实时电价",
            "emergency": "5 × 交易时刻电价",
            "adjustment": "0.5× 下调违约、1.5× 上调违约",
        },
        "result1": audit_result1(panel, checks),
    }
    for tag, file, with_adjust in (("q2", "result2.xlsx", False),
                                   ("q3", "result3.xlsx", True),
                                   ("q4_2", "result4-2.xlsx", False),
                                   ("q4_3", "result4-3.xlsx", True)):
        out[tag] = audit_wide(tag, file, with_adjust, panel, checks)

    failed = [c for c in checks if not c["pass"]]
    out["checks"] = checks
    out["passed"] = len(checks) - len(failed)
    out["failed"] = len(failed)
    out["status"] = "PASS" if not failed else "FAIL"
    RESULTS.mkdir(exist_ok=True)
    (RESULTS / "output_audit.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")

    for c in checks:
        flag = "PASS" if c["pass"] else "FAIL"
        extra = "" if c["max_abs_deviation"] is None else f"  (max dev {c['max_abs_deviation']:.3e})"
        print(f"[{flag}] {c['check']}  {c['detail']}{extra}")
    print(f"\n审计总览：{out['passed']}/{len(checks)} 通过，status={out['status']}")
    print("关键数字：")
    for tag in ("q2", "q3", "q4_2", "q4_3"):
        row = out[tag]
        print(f"  {tag}: 合计 {row['total_cost_cny']:.2f} 元；计划 {row['plan_cost_cny']:.2f} 元；"
              f"紧急 {row['emergency_cost_cny']:.2f} 元 / {row['emergency_energy_kwh']:.2f} kWh；"
              f"违约 {row['over_plan_penalty_cny']:.2f} + {row['under_plan_penalty_cny']:.2f} 元")


if __name__ == "__main__":
    main()
